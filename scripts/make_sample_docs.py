"""Generate the synthetic sample documents used by the demo and the tests.

All content is fictional ("Falcon Ridge Trading LLC"). Several flaws are planted
on purpose so each reliability scenario can be demonstrated:

  * finance_q2_2026_report.docx       operating profit mis-stated (828,700; correct 728,700)
  * board_financial_summary...txt     CONFIDENTIAL; preliminary Q1 revenue differs from final report
  * leave_and_remote_work_policy_2026 annual leave 25 days / remote 3 days vs handbook 22 / 2
  * hr_faq.txt                        probation 3 months vs handbook 6 months, no effective date
  * parental leave                    covered by no document (missing information)
  * bad/corrupt_report.pdf            not a real PDF (processing failure)

Usage:  python scripts/make_sample_docs.py [output_dir]
"""
from __future__ import annotations

import sys
from pathlib import Path

OUT = Path(sys.argv[1] if len(sys.argv) > 1 else Path(__file__).resolve().parent.parent / "sample_docs")


# ---------------------------------------------------------------- helpers
def write_pdf(path: Path, pages: list[list[str]]) -> None:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer

    styles = getSampleStyleSheet()
    story = []
    for i, page in enumerate(pages):
        for line in page:
            if line.startswith("# "):
                story.append(Paragraph(line[2:], styles["Title"]))
            elif line.startswith("## "):
                story.append(Paragraph(line[3:], styles["Heading2"]))
            elif line == "":
                story.append(Spacer(1, 8))
            else:
                story.append(Paragraph(line, styles["BodyText"]))
        if i < len(pages) - 1:
            story.append(PageBreak())
    SimpleDocTemplate(str(path), pagesize=A4, title=path.stem).build(story)


def write_docx(path: Path, blocks: list) -> None:
    import docx

    d = docx.Document()
    for b in blocks:
        if isinstance(b, list):                       # table rows
            t = d.add_table(rows=0, cols=len(b[0]))
            t.style = "Table Grid"
            for row in b:
                cells = t.add_row().cells
                for c, v in zip(cells, row):
                    c.text = v
        elif b.startswith("# "):
            d.add_heading(b[2:], level=1)
        elif b.startswith("## "):
            d.add_heading(b[3:], level=2)
        else:
            d.add_paragraph(b)
    d.save(str(path))


# ---------------------------------------------------------------- documents
def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    (OUT / "bad").mkdir(exist_ok=True)

    write_pdf(OUT / "finance_q1_2026_report.pdf", [[
        "# Q1 2026 Quarterly Financial Report",
        "Reporting period: Q1 2026",
        "Falcon Ridge Trading LLC. Status: final, approved by the CFO on 15 April 2026. All amounts in AED.",
        "",
        "## Income statement",
        "Revenue: AED 4,200,000",
        "Cost of goods sold: AED 2,520,000",
        "Gross profit: AED 1,680,000",
        "Salaries and benefits: AED 620,000",
        "Marketing: AED 180,000",
        "Rent and facilities: AED 150,000",
        "Other operating expenses: AED 90,000",
        "Total operating expenses: AED 1,040,000",
        "Operating profit: AED 640,000",
        "Net profit: AED 512,000",
        "",
        "## Budget and cash",
        "Budgeted revenue: AED 4,000,000",
        "Closing cash balance: AED 2,350,000",
    ], [
        "## Management commentary",
        "Revenue exceeded budget, driven mainly by the new regional distribution contract signed in "
        "February. Freight costs remained elevated and are the main pressure on gross margin.",
        "Marketing spend was kept below plan while the brand refresh was postponed to Q2.",
        "## Risks",
        "Key risks for the coming quarter are fuel price volatility, a concentration of 28 percent of "
        "revenue in two customers, and delayed payments from public-sector clients.",
    ]])

    write_docx(OUT / "finance_q2_2026_report.docx", [
        "# Q2 2026 Quarterly Financial Report",
        "Reporting period: Q2 2026",
        "Falcon Ridge Trading LLC. Status: final, approved by the CFO on 14 July 2026. All amounts in AED.",
        "## Income statement",
        [["Line item", "Q2 2026"],
         ["Revenue", "AED 4,830,000"],
         ["Cost of goods sold", "AED 2,946,300"],
         ["Gross profit", "AED 1,883,700"],
         ["Salaries and benefits", "AED 655,000"],
         ["Marketing", "AED 240,000"],
         ["Rent and facilities", "AED 150,000"],
         ["Other operating expenses", "AED 110,000"],
         ["Total operating expenses", "AED 1,155,000"],
         ["Operating profit", "AED 828,700"],
         ["Net profit", "AED 582,960"]],
        "## Budget and cash",
        [["Item", "Amount"],
         ["Budgeted revenue", "AED 4,600,000"],
         ["Closing cash balance", "AED 2,710,000"]],
        "## Management commentary",
        "Revenue grew on the back of the cold-chain pilot and the full-quarter effect of the "
        "distribution contract. The postponed brand refresh was executed, raising marketing spend.",
        "Cost of goods sold rose faster than revenue because of higher freight rates.",
    ])

    (OUT / "board_financial_summary_confidential.txt").write_text(
        """BOARD OF DIRECTORS - FINANCIAL SUMMARY (CONFIDENTIAL)
Falcon Ridge Trading LLC. Distribution: board members and CFO only.
Reporting period: Q1 2026

Q1 2026 revenue (preliminary): AED 4,150,000

Project Osprey
The board is evaluating the acquisition of Gulf Cold Chain Services for an enterprise value of
AED 12,500,000, funded 60 percent from cash reserves and 40 percent from a bank facility.
Due diligence closes on 30 September 2026. Disclosure of Project Osprey outside the board is prohibited.

Executive compensation
Total CEO compensation approved for 2026 is AED 2,400,000 including the performance bonus.
""")

    write_pdf(OUT / "employee_handbook_2025.pdf", [[
        "# Employee Handbook 2025",
        "Effective date: 1 January 2025. Owner: Human Resources.",
        "## 1. Working hours",
        "Standard working hours are 8 hours per day, Monday to Friday, with core hours between 10:00 and 15:00.",
        "## 2. Probation",
        "New employees serve a probation period of 6 months, during which either party may end employment "
        "with 14 days written notice.",
        "## 3. Notice period",
        "Once employment is confirmed, the notice period is 30 calendar days for grades 1 to 4 and 60 calendar days "
        "for grade 5 and above.",
    ], [
        "## 4. Annual leave",
        "Full-time employees are entitled to 22 working days of paid annual leave per calendar year.",
        "Annual leave requests must be submitted through the HR portal at least 14 days in advance.",
        "## 5. Sick leave",
        "Employees are entitled to 15 days of paid sick leave per year. A medical certificate is required "
        "for absences longer than 2 consecutive days.",
        "## 6. Remote work",
        "Employees may work from home up to 2 days per week with line-manager approval.",
        "## 7. Learning and development",
        "Each employee has an annual training budget of AED 3,000, approved by the line manager.",
        "## 8. Grievances",
        "Grievances are raised in writing with the line manager first, and escalated to HR if unresolved "
        "within 10 working days.",
    ]])

    write_docx(OUT / "leave_and_remote_work_policy_2026.docx", [
        "# Leave and Remote Work Policy 2026",
        "Effective date: 2026-03-01. Owner: Human Resources. This policy supersedes sections 4 and 6 of "
        "the Employee Handbook 2025.",
        "## Annual leave",
        "Full-time employees are entitled to 25 working days of paid annual leave per calendar year.",
        "Up to 5 days of unused annual leave may be carried over to the following year and must be used "
        "by 31 March.",
        "## Remote work",
        "Employees may work remotely up to 3 days per week, subject to team coverage rules.",
        "## Sick leave",
        "Sick leave rules are unchanged from the Employee Handbook.",
    ])

    (OUT / "hr_faq.txt").write_text(
        """HR Frequently Asked Questions
Falcon Ridge Trading LLC - Human Resources

How long is probation for new hires?
The probation period for new hires is 3 months.

Who approves overtime?
Overtime must be pre-approved by the department head and is compensated as time off in lieu.

Where do I submit expense claims?
Expense claims are submitted through the finance portal with original receipts within 30 days.
""")

    (OUT / "salary_bands_2026_confidential.txt").write_text(
        """SALARY BANDS AND BONUS POLICY 2026 (CONFIDENTIAL)
Falcon Ridge Trading LLC - restricted to HR and executive management.
Effective date: 2026-01-01

Monthly base salary bands (AED):
Grade 1: AED 6,000 to AED 9,000
Grade 2: AED 9,000 to AED 14,000
Grade 3: AED 14,000 to AED 21,000
Grade 4: AED 21,000 to AED 32,000
Grade 5: AED 32,000 to AED 48,000

The annual bonus pool is 8 percent of audited net profit, distributed by performance rating.
""")

    (OUT / "company_overview.md").write_text(
        """# Falcon Ridge Trading LLC - Company Overview

Falcon Ridge Trading LLC is a fictional logistics and distribution company founded in 2014 and
headquartered in Dubai, with branch offices in Abu Dhabi and Muscat. It employs about 180 people.

## Services
The company provides contract warehousing, regional distribution and, since 2026, a cold-chain
transport pilot for food and pharmaceutical clients.

## 2026 strategic priorities
1. Scale the cold-chain service to three new clients by the end of 2026.
2. Digitise warehouse operations with barcode scanning and a new warehouse management system.
3. Achieve ISO 27001 information security certification by Q4 2026.

## Leadership
The Chief Executive Officer is Layla Haddad. The Chief Financial Officer is Omar Siddiqui.
""")

    (OUT / "demo_upload").mkdir(exist_ok=True)
    (OUT / "demo_upload" / "marketing_plan_h2_2026.txt").write_text(
        """Marketing Plan H2 2026
Falcon Ridge Trading LLC - Marketing Department
Effective date: 2026-07-01

The second-half marketing plan focuses on the cold-chain service launch and on retaining the two largest
distribution customers.

Budget: the approved marketing budget for H2 2026 is AED 520,000, of which AED 300,000 is digital advertising
and AED 220,000 is trade events.

Decision needed: the board must approve the Riyadh logistics trade show sponsorship of AED 85,000 by 15 August 2026.
""")
    (OUT / "bad" / "corrupt_report.pdf").write_bytes(b"%PDF-1.7\n%garbage that is not a real pdf body\n")
    print(f"sample documents written to {OUT}")


if __name__ == "__main__":
    main()
